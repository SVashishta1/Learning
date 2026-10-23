# make_samples.py: create four realistic, messy Northwind documents to practise on, in the folder samples/.
#
#   python make_samples.py
#
#   samples/delivery_terms.pdf    3 pages: repeated header and footer, words split with hyphens,
#                                 and one page that is only an image (like a scan: no text to extract)
#   samples/rate_card_brouwer.xlsx  a title row, merged cells, blank rows, a notes line under the table, 2 sheets
#   samples/customer_email.eml    an email with plain and HTML versions, a quoted reply chain, a signature
#                                 and a CSV attachment
#   samples/faq.html              a web page full of menus, scripts and a cookie banner around the real content,
#                                 plus a "ligature" (one character for 'fi') pasted from a word processor
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).parent / "samples"

PDF_PAGES = [
    ["3. Delivery",
     "3.1 If a shipment is delayed, Northwind informs the customer of the new expected deliv-",
     "ery time as soon as it is known, by email and in the tracking portal.",
     "3.2 For delays over 24 hours caused by Northwind, business customers may request a 10% credit",
     "on the transport charge. Requests must be made within 30 days of the original delivery date.",
     "3.3 Delays caused by weather, strikes or customs are not eligible for a credit (force majeure)."],
    ["4. Proof of delivery",
     "4.1 Proof of delivery (the name of the person who signed, time and place) is available on re-",
     "quest within 2 working days.",
     "5. Damage",
     "5.1 Damage must be reported within 7 days of delivery. Damage claims are handled by the Claims team,",
     "not Customer Support. Claims require photos and the original packing list."],
]


def make_pdf(path):
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    total = len(PDF_PAGES) + 1
    for number, lines in enumerate(PDF_PAGES, start=1):
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(50, height - 40, "Northwind Logistics - General Delivery Terms v4.2")    # header on every page
        y = height - 90
        for line in lines:
            is_heading = line[0].isdigit() and line[1] == "." and line[2] == " "
            pdf.setFont("Helvetica-Bold" if is_heading else "Helvetica", 12 if is_heading else 10)
            pdf.drawString(50, y, line)
            y -= 22 if is_heading else 16
        pdf.setFont("Helvetica", 8)
        pdf.drawString(50, 30, f"Page {number} of {total} - Confidential")                    # footer on every page
        pdf.showPage()
    # The last page is "scanned": just a picture (here a grey box with lines), no text layer at all.
    pdf.setFillGray(0.85)
    pdf.rect(50, height - 400, width - 100, 320, fill=1, stroke=0)
    pdf.setStrokeGray(0.4)
    for y in range(int(height - 110), int(height - 390), -18):
        pdf.line(70, y, width - 90, y)
    pdf.showPage()
    pdf.save()


def make_excel(path):
    book = Workbook()
    rates = book.active
    rates.title = "Rates"
    rates["A1"] = "Rate card: Brouwer Metals, valid 2026"
    rates["A1"].font = Font(bold=True, size=14)
    rates.merge_cells("A1:D1")                              # a title spread over 4 columns
    rates.append([])                                        # row 2: empty, people like white space
    rates.append(["Zone", "Weight from (kg)", "Weight to (kg)", "Price EUR"])   # row 3: the real header
    rows = [("NL", 0, 100, 45.0), (None, 100, 500, 110.0), (None, 500, 1000, 190.0),
            ("BE / DE", 0, 100, 65.0), (None, 100, 500, 150.0), (None, 500, 1000, 260.0)]
    for row in rows:
        rates.append(list(row))
    rates.merge_cells("A4:A6")                              # "NL" written once for three rows
    rates.merge_cells("A7:A9")
    rates.append([])
    rates.append(["Prices exclude VAT. Fuel surcharge applies, see sheet Surcharges."])
    surcharges = book.create_sheet("Surcharges")
    surcharges.append(["Surcharge", "Amount", "Unit"])
    surcharges.append(["Fuel", 7.5, "% of price"])
    surcharges.append(["Tail-lift delivery", 25, "EUR per shipment"])
    surcharges.append(["Weekend delivery", 40, "EUR per shipment"])
    book.save(path)


def make_email(path):
    msg = EmailMessage()
    msg["From"] = "Joost van Dijk <joost.vandijk@brouwer-metals.example>"
    msg["To"] = "support@northwind.example"
    msg["Subject"] = "RE: Pallet NW-4821907 still not delivered"
    msg["Date"] = "Wed, 21 Oct 2026 09:14:00 +0200"
    msg["Message-ID"] = "<20261021091400.4821907@brouwer-metals.example>"
    plain = ("Hello,\n\n"
             "The pallet NW-4821907 still hasn't arrived and our production line in Utrecht is waiting.\n"
             "Can you confirm the new delivery time today? The list of missing parts is attached.\n\n"
             "Regards,\n"
             "Joost\n"
             "-- \n"
             "Joost van Dijk | Purchasing | Brouwer Metals B.V.\n"
             "+31 30 555 0199 | This email is confidential.\n\n"
             "On Tue, 20 Oct 2026 at 16:02, Northwind Support <support@northwind.example> wrote:\n"
             "> Dear Joost,\n"
             "> Your pallet is on its way and will arrive tomorrow.\n"
             "> Kind regards, Northwind Customer Support\n")
    html = ("<html><body><p>Hello,</p><p>The pallet <b>NW-4821907</b> still hasn't arrived and our production "
            "line in Utrecht is waiting.<br>Can you confirm the new delivery time today? The list of missing parts "
            "is attached.</p><p>Regards,<br>Joost</p></body></html>")
    msg.set_content(plain)
    msg.add_alternative(html, subtype="html")               # the same text again, as HTML
    msg.add_attachment("part_number,description,quantity\nBM-1001,Steel bracket,400\nBM-2040,Hinge,120\n",
                       subtype="csv", filename="missing_parts.csv")
    path.write_bytes(bytes(msg))


FAQ_HTML = """<!doctype html>
<html><head><title>FAQ | Northwind Logistics</title>
<style>body { font-family: sans-serif; }</style>
<script>window.dataLayer = window.dataLayer || []; function track(){ /* analytics */ }</script>
</head>
<body>
<div id="cookie-banner" class="cookie-consent">We use cookies to improve your experience. <button>Accept all</button></div>
<header><nav><a href="/">Home</a> | <a href="/track">Track &amp; trace</a> | <a href="/contact">Contact</a></nav></header>
<main>
  <h1>Frequently asked questions</h1>
  <h2>How do I track my shipment?</h2>
  <p>Enter your tracking number (it looks like NW-1234567) on the Track&nbsp;&amp;&nbsp;trace page.
     Updates appear within 30&nbsp;minutes of each scan.</p>
  <h2>What if my shipment is late?</h2>
  <p>You will receive a new expected delivery time by email. For delays over 24 hours caused by us,
     business customers can request a 10% credit (see the General Delivery Terms, clause 3.2).</p>
  <h2>How do I report damage?</h2>
  <p>Report damage within 7 days of delivery using the claims form. Please include photos and the \ufb01nal packing list.</p>
</main>
<aside class="promo">Ship more, pay less: ask about our volume discounts!</aside>
<footer>&copy; 2026 Northwind Logistics &middot; Privacy &middot; Terms</footer>
<script>track('faq');</script>
</body></html>
"""


def main(out=OUT):
    out = Path(out)
    out.mkdir(exist_ok=True)
    make_pdf(out / "delivery_terms.pdf")
    make_excel(out / "rate_card_brouwer.xlsx")
    make_email(out / "customer_email.eml")
    (out / "faq.html").write_text(FAQ_HTML, encoding="utf-8")
    for path in sorted(out.iterdir()):
        print(f"  {path.name:<24} {path.stat().st_size:>7,} bytes")


if __name__ == "__main__":
    main()
